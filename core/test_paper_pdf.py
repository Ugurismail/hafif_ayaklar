from io import BytesIO
import subprocess
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase
from docx import Document
from pypdf import PdfReader

from core.diagram_markup import encode_diagram_payload
from core.models import Answer, Question, Reference
from core.paper_export import build_paper_docx, _configure_document, _patch_footnotes
from core.paper_pdf import _WordProjection, paper_docx_to_pdf, PDFExportError


class UnifiedDocumentExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(username='document-user', password='pass')
        cls.other = User.objects.create_user(username='other-document-user', password='pass')
        cls.question = Question.objects.create(user=cls.user, question_text='Türkçe belge')
        cls.reference = Reference.objects.create(
            created_by=cls.user, author_name='Ayşe', author_surname='Aksoy',
            year=2020, metin_ismi='Örnek Kaynak', rest='İstanbul.',
        )
        diagram = encode_diagram_payload({
            'title': 'Karar döngüsü',
            'nodes': [{'id': 'a', 'label': 'Başla', 'x': 150, 'y': 200},
                      {'id': 'b', 'label': 'Kontrol', 'x': 600, 'y': 200}],
            'edges': [{'from': 'a', 'to': 'b', 'label': 'Evet', 'description': 'Ok açıklaması.'}],
        })
        cls.answer = Answer.objects.create(user=cls.user, question=cls.question, answer_text=(
            '-- Bölüm Bir\n**Güçlü sav** ve *italik ifade*. '
            f'-g- Sayfa altı açıklaması (k:{cls.reference.id}, s:12). -g-\n\n'
            '---- Alt Bölüm\n[Bağlantı](https://example.com)\n'
            '1.2.3.4.5.6. Derin madde\n\n'
            f'[[diyagram:{diagram}]]'
        ))

    def setUp(self):
        self.client.force_login(self.user)

    def download(self, kind, **kwargs):
        return self.client.post(f'/profile/{self.user.username}/download_entries_{kind}/', {
            'entry_ids': str(self.answer.id), **kwargs,
        })

    def test_word_uses_the_existing_paper_layout_and_legacy_url_is_an_alias(self):
        word = self.download('docx')
        legacy = self.download('paper')
        self.assertEqual(word.status_code, 200)
        self.assertIn('_entries.docx', word['Content-Disposition'])
        doc = Document(BytesIO(word.content))
        self.assertEqual(doc.styles['Paper Body'].font.name, 'Garamond')
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertEqual([p.text for p in doc.paragraphs],
                         [p.text for p in Document(BytesIO(legacy.content)).paragraphs])
        self.assertNotEqual(doc.core_properties.title, 'Paper')

    def test_pdf_preserves_hierarchy_notes_bibliography_diagrams_and_links(self):
        response = self.download('pdf')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        reader = PdfReader(BytesIO(response.content))
        text = '\n'.join(page.extract_text() for page in reader.pages)
        for content in ('İçindekiler', 'Türkçe belge', 'Bölüm Bir', 'Alt Bölüm',
                        'Güçlü sav', 'italik ifade', 'Derin madde', 'Sayfa altı açıklaması',
                        'Kaynakça', 'Örnek Kaynak', 'Diyagram 1.', 'Ok açıklaması'):
            self.assertIn(content, text)
        for raw_markup in ('[[diyagram:', '-g-', '**Güçlü', '(k:'):
            self.assertNotIn(raw_markup, text)
        self.assertTrue(reader.outline)
        self.assertTrue(any(page.images for page in reader.pages))
        annotations = [item.get_object() for page in reader.pages for item in page.get('/Annots', [])]
        self.assertTrue(any(item.get('/A', {}).get('/URI') == 'https://example.com' for item in annotations))
        self.assertTrue(any(item.get('/Dest') for item in annotations))

    def test_selection_and_custom_order_are_shared_by_both_formats(self):
        second = Answer.objects.create(user=self.user, question=self.question, answer_text='Önce seçilen metin.')
        Answer.objects.create(user=self.other, question=self.question, answer_text='Başkasının özel seçimi.')
        for kind in ('docx', 'pdf'):
            response = self.download(kind, entry_ids=f'{second.id},{self.answer.id}', order='custom')
            if kind == 'pdf':
                text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)
            else:
                text = '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
            self.assertLess(text.index('Önce seçilen metin.'), text.index('Güçlü sav'))
            self.assertNotIn('Başkasının özel seçimi', text)

    def test_both_formats_reject_other_users_before_rendering(self):
        self.client.force_login(self.other)
        with patch('core.paper_export.build_paper_docx') as build:
            for kind in ('docx', 'pdf', 'paper'):
                self.assertEqual(self.download(kind).status_code, 403)
        build.assert_not_called()

    def test_projection_escapes_content_and_refuses_external_and_local_resources(self):
        self.answer.answer_text = '<img src="file:///etc/passwd"><script>alert(1)</script>\n[jump](javascript:alert)'
        projection = _WordProjection(build_paper_docx([self.answer], self.user))
        html = projection.html()
        self.assertNotIn('<script>', html)
        self.assertNotIn('href="javascript:', html)
        self.assertNotIn('<img src="file:', html)
        for url in ('file:///etc/passwd', 'http://127.0.0.1/private', 'https://example.com/a.png', 'data:text/html,test'):
            with self.assertRaises(ValueError):
                projection.fetch_asset(url)

    def test_long_note_is_not_lost_at_a_page_boundary(self):
        self.answer.answer_text = ('Uzun paragraf. ' * 260 + '\n\nSon sav '
                                   '-g- ' + 'Ayrıntılı dipnot. ' * 100 + ' NOTUN SONU -g-')
        reader = PdfReader(BytesIO(paper_docx_to_pdf(build_paper_docx([self.answer], self.user))))
        self.assertIn('NOTUN SONU', '\n'.join(p.extract_text() for p in reader.pages))

    def test_footnote_marks_restart_on_each_page_and_stay_with_the_call(self):
        document = Document()
        _configure_document(document)
        document.add_paragraph('Birinci sav [[PAPER_FOOTNOTE_1]]', style='Paper Body')
        document.add_page_break()
        document.add_paragraph('İkinci sav [[PAPER_FOOTNOTE_2]]', style='Paper Body')
        source = BytesIO()
        document.save(source)
        data = _patch_footnotes(source.getvalue(), [
            {'id': 1, 'marker': '[[PAPER_FOOTNOTE_1]]', 'text': 'İlk sayfanın notu.'},
            {'id': 2, 'marker': '[[PAPER_FOOTNOTE_2]]', 'text': 'İkinci sayfanın notu.'},
        ])
        pages = PdfReader(BytesIO(paper_docx_to_pdf(data))).pages
        self.assertEqual(len(pages), 2, [page.extract_text() for page in pages])
        for page, expected in zip(pages, ['İlk sayfanın notu.', 'İkinci sayfanın notu.']):
            text = page.extract_text()
            self.assertIn(expected, text)
            self.assertIn('*', text)
            self.assertNotIn('†', text)

    def test_pdf_worker_timeout_is_bounded_and_word_remains_available(self):
        with patch('core.paper_pdf.subprocess.run', side_effect=subprocess.TimeoutExpired('pdf', 60)):
            response = self.download('pdf')
        self.assertEqual(response.status_code, 503)
        self.assertContains(response, 'süresi aşıldı', status_code=503)
        self.assertEqual(self.download('docx').status_code, 200)

    def test_pdf_worker_rejects_non_pdf_output(self):
        with patch('core.paper_pdf.subprocess.run', return_value=subprocess.CompletedProcess([], 0, b'not a pdf')):
            with self.assertRaises(PDFExportError):
                paper_docx_to_pdf(b'test')
